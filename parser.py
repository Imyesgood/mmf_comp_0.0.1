import os, re, zipfile, tempfile
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.utils import column_index_from_string
from config import FUND_MAP, FUND_ORDER, AUTO_ROWS


def normalize(s) -> str:
    if not s: return ""
    s = str(s)
    s = re.sub(r"[\(\[【][^\)\]】]*[\)\]】]", "", s)
    s = re.sub(r"[-_]운용$", "", s)
    return re.sub(r"\s+", "", s).strip()


def sanitize_named_styles(in_path):
    out = os.path.join(tempfile.gettempdir(), os.path.basename(in_path) + "_san.xlsx")
    with tempfile.TemporaryDirectory() as td:
        with zipfile.ZipFile(in_path) as z: z.extractall(td)
        sp = os.path.join(td, "xl", "styles.xml")
        if os.path.exists(sp):
            tree = ET.parse(sp); root = tree.getroot()
            for i, el in enumerate(e for e in root.iter() if e.tag.endswith("cellStyle")):
                if el.get("name") in (None, ""): el.set("name", f"S{i}")
            tree.write(sp, encoding="utf-8", xml_declaration=True)
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            for fd, _, fs in os.walk(td):
                for f in fs:
                    fp = os.path.join(fd, f); z.write(fp, os.path.relpath(fp, td))
    return out


def load_wb(path, data_only=False):
    try: return openpyxl.load_workbook(path, data_only=data_only)
    except: return openpyxl.load_workbook(sanitize_named_styles(path), data_only=data_only)


def parse_mm_my(path, sheet):
    wb = load_wb(path, data_only=True); ws = wb[sheet]
    code_to_row, file_names = {}, {}
    for r in range(2, ws.max_row + 1):
        code = ws.cell(r, 1).value; name = ws.cell(r, 2).value
        if code and str(code).strip():
            code = str(code).strip()
            code_to_row[code] = r
            file_names[code] = str(name).strip() if name else ""
    wb.close()
    for code, expected in FUND_MAP.items():
        if code not in code_to_row:
            raise ValueError(f"[파서 오류] {os.path.basename(path)}: 펀드코드 {code} ({expected}) 없음")
        actual = normalize(file_names[code])
        if actual != expected:
            raise ValueError(
                f"[파서 오류] {os.path.basename(path)}: {code} 펀드명 불일치\n"
                f"  FUND_MAP: {expected!r}\n  파일: {file_names[code]!r} (정규화: {actual!r})\n"
                f"  → FUND_MAP['{code}'] = {actual!r} 로 수정 또는 normalize() 수정"
            )
    return code_to_row


def parse_f2(path, sheet):
    wb = load_wb(path, data_only=True); ws = wb[sheet]
    norm_to_row, norm_to_raw = {}, {}
    for r in range(3, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if name and str(name).strip():
            n = normalize(str(name)); norm_to_row[n] = r; norm_to_raw[n] = str(name).strip()
    wb.close()
    code_to_row = {}
    for code, expected in FUND_MAP.items():
        n = normalize(expected)
        if n not in norm_to_row:
            raise ValueError(
                f"[파서 오류] F2: {code} ({expected}) 없음\n"
                f"  F2 목록:\n" + "\n".join(f"    {v!r}" for v in norm_to_raw.values()) +
                f"\n  → FUND_MAP['{code}'] 수정 또는 normalize() 수정"
            )
        code_to_row[code] = norm_to_row[n]
    return code_to_row


def _fmt(val, fmt_type):
    if val is None or (isinstance(val, str) and not val.strip()): return "-"
    if isinstance(val, (int, float)):
        if fmt_type == "float2": return f"{float(val):.2f}"
        if fmt_type == "int_comma": return f"{int(round(float(val))):,}"
    return str(val)


def build_table(mm_path, my_path, f2_path, cmp_path=None):
    """반환: { row_label: { fund_code: 값(str) } }  (AUTO_ROWS만)"""
    mm_rows = parse_mm_my(mm_path, "FunddoctorPro V2_output")
    my_rows = parse_mm_my(my_path, "FunddoctorPro V2_output")
    f2_rows = parse_f2(f2_path, "Sheet1")

    mm_wb = load_wb(mm_path, data_only=True)
    my_wb = load_wb(my_path, data_only=True)
    f2_wb = load_wb(f2_path, data_only=True)
    mm_ws = mm_wb["FunddoctorPro V2_output"]
    my_ws = my_wb["FunddoctorPro V2_output"]
    f2_ws = f2_wb["Sheet1"]

    tag_map = {
        "F^2_H":  (f2_ws, f2_rows,  "H"),
        "M_M_E":  (mm_ws, mm_rows, "E"),
        "M_Y_E":  (my_ws, my_rows, "E"),
        "F^2_T":  (f2_ws, f2_rows,  "T"),
        "F^2_W":  (f2_ws, f2_rows,  "W"),
        "F^2_AC": (f2_ws, f2_rows,  "AC"),
    }

    result = {}
    for label, cmp_code, fmt_type in AUTO_ROWS:
        src_ws, row_map, col_letter = tag_map[cmp_code]
        col_idx = column_index_from_string(col_letter)
        result[label] = {}
        for code in FUND_ORDER:
            sr = row_map.get(code)
            val = src_ws.cell(sr, col_idx).value if sr else None
            result[label][code] = _fmt(val, fmt_type)

    for wb in (mm_wb, my_wb, f2_wb): wb.close()
    return result
